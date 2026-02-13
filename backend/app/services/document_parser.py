"""
Document Parser Service - Multi-format document parsing

Supports .txt, .md, .pdf, and .docx file parsing with a unified interface.
"""
from typing import Optional
import logging

logger = logging.getLogger(__name__)


class DocumentParser:
    """Parses various document formats into plain text."""
    
    SUPPORTED_EXTENSIONS = {".txt", ".md", ".text", ".pdf", ".docx"}
    
    def parse(self, file_bytes: bytes, filename: str) -> str:
        """
        Parse a file into plain text.
        
        Args:
            file_bytes: Raw file bytes
            filename: Original filename (used for extension detection)
            
        Returns:
            Extracted text content
            
        Raises:
            ValueError: If file type is unsupported
        """
        ext = self._get_extension(filename)
        
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type '{ext}'. "
                f"Supported: {', '.join(sorted(self.SUPPORTED_EXTENSIONS))}"
            )
        
        if ext in {".txt", ".md", ".text"}:
            return self._parse_text(file_bytes)
        elif ext == ".pdf":
            return self._parse_pdf(file_bytes)
        elif ext == ".docx":
            return self._parse_docx(file_bytes)
        
        return self._parse_text(file_bytes)
    
    def _get_extension(self, filename: str) -> str:
        """Extract file extension."""
        if "." in filename:
            return "." + filename.rsplit(".", 1)[-1].lower()
        return ""
    
    def _parse_text(self, file_bytes: bytes) -> str:
        """Parse plain text / markdown files."""
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")
    
    def _parse_pdf(self, file_bytes: bytes) -> str:
        """Parse PDF files using PyPDF2."""
        try:
            from PyPDF2 import PdfReader
            import io
            
            reader = PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(f"--- Page {i + 1} ---\n{page_text}")
            
            if not text_parts:
                logger.warning("PDF parsing returned no text. The PDF may be image-based.")
                return "[PDF contained no extractable text. It may be scanned/image-based.]"
            
            return "\n\n".join(text_parts)
            
        except ImportError:
            logger.error("PyPDF2 not installed. Run: pip install PyPDF2")
            raise ValueError("PDF support requires PyPDF2. Install with: pip install PyPDF2")
        except Exception as e:
            logger.error(f"PDF parsing failed: {e}")
            raise ValueError(f"Failed to parse PDF: {str(e)}")
    
    def _parse_docx(self, file_bytes: bytes) -> str:
        """Parse DOCX files using python-docx."""
        try:
            from docx import Document
            import io
            
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # Preserve heading structure
                    if para.style and para.style.name.startswith("Heading"):
                        level = para.style.name.replace("Heading ", "")
                        try:
                            hashes = "#" * int(level)
                            paragraphs.append(f"{hashes} {text}")
                        except ValueError:
                            paragraphs.append(f"## {text}")
                    else:
                        paragraphs.append(text)
            
            # Also extract tables
            for table in doc.tables:
                table_text = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    table_text.append(" | ".join(cells))
                if table_text:
                    paragraphs.append("\n".join(table_text))
            
            if not paragraphs:
                return "[DOCX contained no extractable text.]"
            
            return "\n\n".join(paragraphs)
            
        except ImportError:
            logger.error("python-docx not installed. Run: pip install python-docx")
            raise ValueError("DOCX support requires python-docx. Install with: pip install python-docx")
        except Exception as e:
            logger.error(f"DOCX parsing failed: {e}")
            raise ValueError(f"Failed to parse DOCX: {str(e)}")
    
    def is_supported(self, filename: str) -> bool:
        """Check if a file type is supported."""
        return self._get_extension(filename) in self.SUPPORTED_EXTENSIONS


# Singleton instance
document_parser = DocumentParser()
